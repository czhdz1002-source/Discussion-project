"""
三表匹配数据转移脚本 v4
════════════════════════════════════════════════════════════
源数据表标记与匹配键：
  新三板利润表    → A(列1=证券代码) + B(列3=报告期1)
  新三板资产负债表 → A(列1=证券代码) + C(列3=报告期2)
  新三板现金流量表 → A(列1=证券代码) + D(列3=报告期3)

数据处理样板（Sheet1）区域：
  第0行标记：A=列1, B=列9
  第2行子表头：报告期1=列9, 报告期2=列61, 报告期3=列172
  填充区域：利润表(列9~59) | 资产负债表(列61~170) | 现金流量表(列172~283)

匹配逻辑：以(证券代码, 报告期) 为 key，按字段名对齐写入
  标记字母 B/C/D 决定源表中报告期所在列，与样板无关
  样板中报告期列通过子表头字段名"报告期1/2/3"自动定位
════════════════════════════════════════════════════════════
"""

import pandas as pd
from openpyxl import load_workbook

# ===================== 请在此处填写文件路径 =====================
profit_path   = r"D:\This computer\desktop\新三板利润表.xlsx"
bs_path       = r"D:\This computer\desktop\新三板资产负债表.xlsx"
cf_path       = r"D:\This computer\desktop\新三板现金流量表.xlsx"
template_path = r"D:\This computer\desktop\数据处理样板（1）.xlsx"
output_path   = r"D:\This computer\desktop\数据处理\新建 XLSX 工作表.xlsx"
# ==============================================================


def build_source_map(filepath, match_mark, sheet_name=0):
    """
    读取源数据表，构建匹配字典。
    
    参数:
      match_mark: 报告期列的标记字母，利润表='B'，资产负债表='C'，现金流量表='D'
    
    返回:
      source_map: {(证券代码, 报告期): {字段名: 值}}
      col_names:  报告期列之后的所有字段名列表
    """
    df = pd.read_excel(filepath, sheet_name=sheet_name, header=None, dtype=str)

    # 第0行：A标记=证券代码列，match_mark标记=报告期列
    a_col = df.iloc[0][df.iloc[0] == 'A'].index[0]
    b_col = df.iloc[0][df.iloc[0] == match_mark].index[0]

    # 第1行：真正表头；报告期列及之后的字段为待转移数据
    header    = df.iloc[1]
    data_cols = [i for i in range(b_col, len(df.columns)) if pd.notna(header[i])]
    col_names = [header[i] for i in data_cols]

    # 第2行起：实际数据，过滤 A列或报告期列为空的行
    df_data = df.iloc[2:].copy()
    df_data = df_data[df_data[a_col].notna() & df_data[b_col].notna()]

    source_map = {}
    for _, row in df_data.iterrows():
        key = (str(row[a_col]).strip(), str(row[b_col]).strip())
        source_map[key] = {
            col_names[i]: row[data_cols[i]] for i in range(len(data_cols))
        }

    return source_map, col_names


def build_template_col_map(df_template, col_start, col_end, valid_fields):
    """从样板第2行子表头提取 {字段名: 列索引}，仅保留 valid_fields 中存在的字段。"""
    sub = df_template.iloc[2]
    return {
        sub[i]: i
        for i in range(col_start, col_end + 1)
        if i < len(sub) and pd.notna(sub[i]) and sub[i] in valid_fields
    }


def match_and_fill(ws, df_template, source_map, template_col_map,
                   tmpl_a_col, tmpl_period_col, label):
    """遍历样板数据行，匹配后逐字段写入 openpyxl worksheet。"""
    matched = skip = 0
    for row_idx in range(3, len(df_template)):
        row = df_template.iloc[row_idx]
        # 跳过证券代码或报告期为空的行
        if pd.isna(row[tmpl_a_col]) or pd.isna(row[tmpl_period_col]):
            skip += 1
            continue
        key = (str(row[tmpl_a_col]).strip(), str(row[tmpl_period_col]).strip())
        if key in source_map:
            excel_row = row_idx + 1  # pandas行索引→Excel行号(1-based)
            for field_name, col_idx in template_col_map.items():
                value = source_map[key].get(field_name)
                if pd.notna(value) and str(value).strip() not in ('', 'nan'):
                    ws.cell(row=excel_row, column=col_idx + 1, value=value)
            matched += 1

    miss = len(df_template) - 3 - skip - matched
    print(f"  [{label}] 匹配成功: {matched} 行 | 跳过空键: {skip} 行 | 未命中: {miss} 行")
    return matched


# ════════════════════ 主流程 ════════════════════════════════

print("=" * 58)
print("Step 1: 读取三张源数据表...")

#   利润表    → 报告期标记 B
#   资产负债表 → 报告期标记 C
#   现金流量表 → 报告期标记 D
profit_map, profit_fields = build_source_map(profit_path, match_mark='B')
bs_map,     bs_fields     = build_source_map(bs_path,     match_mark='C')
cf_map,     cf_fields     = build_source_map(cf_path,     match_mark='D')

print(f"  利润表    有效数据行: {len(profit_map):>6}  字段数: {len(profit_fields)}")
print(f"  资产负债表 有效数据行: {len(bs_map):>6}  字段数: {len(bs_fields)}")
print(f"  现金流量表 有效数据行: {len(cf_map):>6}  字段数: {len(cf_fields)}")

print("\nStep 2: 读取数据处理样板...")
df_template = pd.read_excel(template_path, sheet_name="Sheet1",
                             header=None, dtype=str)
print(f"  样板数据行: {len(df_template)-3} 行  总列数: {len(df_template.columns)}")

# 样板证券代码列（A标记）
tmpl_a_col = df_template.iloc[0][df_template.iloc[0] == 'A'].index[0]

# 样板各报告期列：通过子表头字段名自动定位
sub = df_template.iloc[2]
profit_period_col = next(i for i in range(9,   60)  if pd.notna(sub[i]) and sub[i] == '报告期1')
bs_period_col     = next(i for i in range(61,  172) if pd.notna(sub[i]) and sub[i] == '报告期2')
cf_period_col     = next(i for i in range(172, 284) if pd.notna(sub[i]) and sub[i] == '报告期3')

print(f"  A=列{tmpl_a_col}(证券代码) | "
      f"报告期1=列{profit_period_col} | "
      f"报告期2=列{bs_period_col} | "
      f"报告期3=列{cf_period_col}")

# 各区域字段名→样板列索引 映射
profit_col_map = build_template_col_map(df_template, 9,   59,  set(profit_fields))
bs_col_map     = build_template_col_map(df_template, 61,  170, set(bs_fields))
cf_col_map     = build_template_col_map(df_template, 172, 283, set(cf_fields))

print(f"  可填充字段: 利润表 {len(profit_col_map)} | "
      f"资产负债表 {len(bs_col_map)} | "
      f"现金流量表 {len(cf_col_map)}")

print("\nStep 3: 匹配并写入数据...")
wb = load_workbook(template_path)
ws = wb["Sheet1"]

match_and_fill(ws, df_template, profit_map, profit_col_map,
               tmpl_a_col, profit_period_col, "利润表   ")
match_and_fill(ws, df_template, bs_map,     bs_col_map,
               tmpl_a_col, bs_period_col,     "资产负债表")
match_and_fill(ws, df_template, cf_map,     cf_col_map,
               tmpl_a_col, cf_period_col,     "现金流量表")

print("\nStep 4: 保存结果...")
wb.save(output_path)
print(f"\n✅ 完成！结果已保存至：{output_path}")
print("=" * 58)
